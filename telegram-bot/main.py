
import logging
import sqlite3
import csv
import os
import asyncio
import hashlib
import smtplib
from contextlib import asynccontextmanager
from datetime import datetime
from email.message import EmailMessage
from html import escape
from http import HTTPStatus
from zoneinfo import ZoneInfo

import uvicorn
from starlette.applications import Starlette
from starlette.requests import Request
from starlette.responses import PlainTextResponse, Response
from starlette.routing import Route
from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import Application, CommandHandler, MessageHandler, CallbackQueryHandler, ContextTypes, filters
from config import DB_PATH, EXPORT_DIR
from database_setup import init_db

TOKEN = os.getenv("BOT_TOKEN")
SMTP_HOST = os.getenv("SMTP_HOST")
SMTP_PORT = int(os.getenv("SMTP_PORT", "587"))
SMTP_USERNAME = os.getenv("SMTP_USERNAME")
SMTP_PASSWORD = os.getenv("SMTP_PASSWORD")
SMTP_FROM_EMAIL = os.getenv("SMTP_FROM_EMAIL")
SMTP_USE_TLS = os.getenv("SMTP_USE_TLS", "true").lower() == "true"
APP_TIMEZONE = os.getenv("APP_TIMEZONE", "Asia/Yangon")
MONTHLY_REPORT_RECIPIENT_EMAIL = os.getenv("MONTHLY_REPORT_RECIPIENT_EMAIL") or SMTP_FROM_EMAIL
MONTHLY_REPORT_SEND_HOUR = int(os.getenv("MONTHLY_REPORT_SEND_HOUR", "9"))
MONTHLY_REPORT_SEND_MINUTE = int(os.getenv("MONTHLY_REPORT_SEND_MINUTE", "0"))

if not TOKEN:
    raise RuntimeError(
        "BOT_TOKEN is missing. Create a .env file next to main.py and set BOT_TOKEN=your_telegram_bot_token"
    )

# Enable logging (Logging ဖွင့်ခြင်း)
logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s", level=logging.INFO
)
logger = logging.getLogger(__name__)
LOCAL_TZ = ZoneInfo(APP_TIMEZONE)


def normalize_url_path(raw_path: str) -> str:
    path = raw_path.strip() or "/"
    return path if path.startswith("/") else f"/{path}"


def env_flag(name: str, default: bool = False) -> bool:
    value = os.getenv(name)
    if value is None:
        return default
    return value.strip().lower() in {"1", "true", "yes", "on"}


def build_webhook_url(base_url: str, path: str) -> str:
    return f"{base_url.rstrip('/')}{path}"


PORT = int(os.getenv("PORT", "10000"))
WEBHOOK_BASE_URL = os.getenv("WEBHOOK_BASE_URL") or os.getenv("RENDER_EXTERNAL_URL")
WEBHOOK_PATH = normalize_url_path(os.getenv("WEBHOOK_PATH", "/telegram/webhook"))
HEALTHCHECK_PATH = normalize_url_path(os.getenv("HEALTHCHECK_PATH", "/healthz"))
WEBHOOK_SECRET_TOKEN = os.getenv("WEBHOOK_SECRET_TOKEN") or hashlib.sha256(TOKEN.encode("utf-8")).hexdigest()
WEBHOOK_LISTEN_HOST = os.getenv("WEBHOOK_LISTEN_HOST", "0.0.0.0")
DROP_PENDING_UPDATES = env_flag("DROP_PENDING_UPDATES", default=False)
BOT_RUN_MODE = os.getenv(
    "BOT_RUN_MODE",
    "webhook" if WEBHOOK_BASE_URL or os.getenv("RENDER_SERVICE_TYPE") == "web" else "polling",
).strip().lower()
FREE_PLAN_MODE = env_flag("FREE_PLAN_MODE", default=False)
EMAIL_DELIVERY_ENABLED = env_flag("EMAIL_DELIVERY_ENABLED", default=not FREE_PLAN_MODE)
MONTHLY_REPORT_ENABLED = env_flag("MONTHLY_REPORT_ENABLED", default=not FREE_PLAN_MODE)
USER_EMAIL_REQUIRED = env_flag("USER_EMAIL_REQUIRED", default=EMAIL_DELIVERY_ENABLED)

SERVICE_OPTIONS = [
    "Body Service",
    "Windows Service",
    "Device Clean",
    "Software Installation",
    "Small Repairs",
]

# Database Functions (ဒေတာဘေ့စ် လုပ်ဆောင်ချက်များ)
def get_db_connection():
    conn = sqlite3.connect(str(DB_PATH))
    conn.row_factory = sqlite3.Row  # Dictionary-like access (အဘိဓာန်ကဲ့သို့ ဝင်ရောက်နိုင်ရန်)
    return conn

def add_service_log(technician_name, device_id, user_name, user_email, service_done, parts_used, technician_chat_id):
    delivery_method = "email" if EMAIL_DELIVERY_ENABLED and user_email else "telegram"
    status = "Email Pending" if delivery_method == "email" else "Saved Locally"
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        INSERT INTO service_logs (
            technician_name, device_id, user_name, user_email,
            service_done, parts_used, technician_chat_id, delivery_method, status
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            technician_name,
            device_id,
            user_name,
            user_email,
            service_done,
            parts_used,
            technician_chat_id,
            delivery_method,
            status,
        ),
    )
    log_id = cursor.lastrowid
    conn.commit()
    conn.close()
    return log_id

def get_service_log(log_id):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM service_logs WHERE id = ?", (log_id,))
    log = cursor.fetchone()
    conn.close()
    return log

def update_service_log_status(log_id, status):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("UPDATE service_logs SET status = ? WHERE id = ?", (status, log_id))
    conn.commit()
    conn.close()

def add_confirmation(service_log_id, user_id, user_telegram_name):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute(
        "INSERT INTO confirmations (service_log_id, user_id, user_telegram_name) VALUES (?, ?, ?)",
        (service_log_id, user_id, user_telegram_name),
    )
    conn.commit()
    conn.close()

def get_technician_service_logs(chat_id):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM service_logs WHERE technician_chat_id = ? ORDER BY created_at DESC", (chat_id,))
    logs = cursor.fetchall()
    conn.close()
    return logs

def get_all_service_logs():
    """Service log အားလုံးကို ရယူခြင်း (Export အတွက်)"""
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT s.id, s.technician_name, s.device_id, s.user_name, s.user_email,
               s.service_done, s.parts_used, s.status, s.created_at,
               c.user_telegram_name as confirmed_by, c.confirmed_at
        FROM service_logs s
        LEFT JOIN confirmations c ON s.id = c.service_log_id
        ORDER BY s.created_at DESC
    """)
    logs = cursor.fetchall()
    conn.close()
    return logs


def get_setting(key: str):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT value FROM app_settings WHERE key = ?", (key,))
    row = cursor.fetchone()
    conn.close()
    return row["value"] if row else None


def set_setting(key: str, value: str) -> None:
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        INSERT INTO app_settings (key, value)
        VALUES (?, ?)
        ON CONFLICT(key) DO UPDATE SET value = excluded.value
        """,
        (key, value),
    )
    conn.commit()
    conn.close()

# Export Functions (Export လုပ်ဆောင်ချက်များ)
def export_to_csv():
    """Service log များကို CSV ဖိုင်အဖြစ် export ထုတ်ခြင်း"""
    logs = get_all_service_logs()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = EXPORT_DIR / f"service_logs_{timestamp}.csv"
    
    with open(filename, 'w', newline='', encoding='utf-8') as csvfile:
        writer = csv.writer(csvfile)
        # Header row (ခေါင်းစီးတန်း)
        writer.writerow([
            'Log ID', 'Technician Name', 'Device ID', 'User Name', 'User Email',
            'Service Done', 'Parts Used', 'Status', 'Created Date',
            'Confirmed By', 'Confirmed Date'
        ])
        # Data rows (ဒေတာတန်းများ)
        for log in logs:
            writer.writerow([
                log['id'], log['technician_name'], log['device_id'],
                log['user_name'], log['user_email'] or 'N/A', log['service_done'], log['parts_used'],
                log['status'], log['created_at'],
                log['confirmed_by'] or 'N/A', log['confirmed_at'] or 'N/A'
            ])
    
    return str(filename)

def export_to_excel():
    """Service log များကို Excel ဖိုင်အဖြစ် export ထုတ်ခြင်း"""
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    except ImportError:
        return None  # openpyxl မရှိပါက None ပြန်ပေးခြင်း
    
    logs = get_all_service_logs()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = EXPORT_DIR / f"service_logs_{timestamp}.xlsx"
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Service Logs"
    
    # Header style (ခေါင်းစီး ပုံစံ)
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Headers (ခေါင်းစီးများ)
    headers = [
        'Log ID', 'Technician Name', 'Device ID', 'User Name', 'User Email',
        'Service Done', 'Parts Used', 'Status', 'Created Date',
        'Confirmed By', 'Confirmed Date'
    ]
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    # Data rows (ဒေတာတန်းများ)
    confirmed_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    pending_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    
    for row_idx, log in enumerate(logs, 2):
        row_data = [
            log['id'], log['technician_name'], log['device_id'],
            log['user_name'], log['user_email'] or 'N/A', log['service_done'], log['parts_used'],
            log['status'], log['created_at'],
            log['confirmed_by'] or 'N/A', log['confirmed_at'] or 'N/A'
        ]
        
        # Status အလိုက် အရောင်ခြယ်ခြင်း
        row_fill = confirmed_fill if log['status'] == 'Confirmed' else pending_fill
        
        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.border = thin_border
            cell.fill = row_fill
    
    # Column width ချိန်ညှိခြင်း
    column_widths = [8, 18, 15, 18, 28, 35, 25, 20, 20, 18, 20]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    
    wb.save(filename)
    return str(filename)

# State management for new service log creation
user_data = {}


def email_config_is_ready() -> bool:
    return not get_missing_smtp_settings()


def get_missing_smtp_settings() -> list[str]:
    missing = []
    for key, value in {
        "SMTP_HOST": SMTP_HOST,
        "SMTP_PORT": SMTP_PORT,
        "SMTP_USERNAME": SMTP_USERNAME,
        "SMTP_PASSWORD": SMTP_PASSWORD,
        "SMTP_FROM_EMAIL": SMTP_FROM_EMAIL,
    }.items():
        if value in (None, ""):
            missing.append(key)
    return missing


def get_placeholder_smtp_settings() -> list[str]:
    placeholders = []
    values = {
        "SMTP_USERNAME": SMTP_USERNAME,
        "SMTP_PASSWORD": SMTP_PASSWORD,
        "SMTP_FROM_EMAIL": SMTP_FROM_EMAIL,
    }
    placeholder_markers = (
        "your_",
        "example@",
        "password",
        "app_password",
    )

    for key, value in values.items():
        if isinstance(value, str) and any(marker in value.lower() for marker in placeholder_markers):
            placeholders.append(key)
    return placeholders


def is_valid_email(email: str) -> bool:
    if "@" not in email:
        return False
    local_part, _, domain = email.partition("@")
    return bool(local_part and domain and "." in domain)


def describe_email_error(exc: Exception) -> str:
    if isinstance(exc, smtplib.SMTPAuthenticationError):
        return (
            "SMTP login failed. Microsoft account password/app password may be wrong, "
            "or SMTP AUTH may be disabled on the mailbox."
        )
    if isinstance(exc, smtplib.SMTPConnectError):
        return "SMTP server connection failed. Check SMTP_HOST, SMTP_PORT, and internet access."
    if isinstance(exc, smtplib.SMTPServerDisconnected):
        return "SMTP server disconnected unexpectedly. TLS or server policy may be blocking the login."
    if isinstance(exc, smtplib.SMTPRecipientsRefused):
        return "Recipient email was refused by the SMTP server. Check the user's email address."
    if isinstance(exc, smtplib.SMTPSenderRefused):
        return "Sender email was refused. Check SMTP_FROM_EMAIL and mailbox permissions."
    if isinstance(exc, smtplib.SMTPDataError):
        return "SMTP server rejected the message content or sender policy."
    if isinstance(exc, TimeoutError):
        return "Connection timed out while contacting the SMTP server."
    if isinstance(exc, OSError):
        return "Network or socket error while connecting to the SMTP server."
    return "Unexpected SMTP error occurred."


def build_service_slip_text(
    user_name: str,
    device_id: str,
    technician_name: str,
    service_done: str,
    parts_used: str,
    service_date: str,
) -> str:
    return (
        "IT Service Slip\n\n"
        f"Name : {user_name}\n"
        f"Device Name : {device_id}\n"
        f"Service By : {technician_name}\n"
        f"Service : {service_done}\n"
        f"Parts Change : {parts_used}\n"
        f"Date : {service_date}\n"
    )


def build_service_slip_html(
    user_name: str,
    device_id: str,
    technician_name: str,
    service_done: str,
    parts_used: str,
    service_date: str,
) -> str:
    fields = [
        ("Name", user_name),
        ("Device Name", device_id),
        ("Service By", technician_name),
        ("Service", service_done),
        ("Parts Change", parts_used),
        ("Date", service_date),
    ]
    rows = "".join(
        (
            "<tr>"
            f"<td style=\"padding:12px 16px;font-weight:700;color:#0f172a;width:170px;border-bottom:1px solid #e2e8f0;\">{escape(label)}</td>"
            f"<td style=\"padding:12px 16px;color:#334155;border-bottom:1px solid #e2e8f0;\">{escape(value)}</td>"
            "</tr>"
        )
        for label, value in fields
    )
    return (
        "<html>"
        "<body style=\"margin:0;padding:24px;background:#f8fafc;font-family:Arial,sans-serif;color:#0f172a;\">"
        "<div style=\"max-width:640px;margin:0 auto;background:#ffffff;border:1px solid #cbd5e1;border-radius:18px;overflow:hidden;box-shadow:0 10px 30px rgba(15,23,42,0.08);\">"
        "<div style=\"padding:24px 28px;background:linear-gradient(135deg,#1d4ed8,#0f766e);color:#ffffff;\">"
        "<div style=\"font-size:12px;letter-spacing:1.5px;text-transform:uppercase;opacity:0.9;\">IT Service Bot</div>"
        "<h2 style=\"margin:8px 0 0;font-size:28px;\">Service Slip</h2>"
        "</div>"
        "<div style=\"padding:24px 28px;\">"
        "<p style=\"margin:0 0 18px;font-size:15px;color:#475569;\">Your device service record is ready.</p>"
        "<table style=\"width:100%;border-collapse:collapse;border:1px solid #e2e8f0;border-radius:12px;overflow:hidden;\">"
        f"{rows}"
        "</table>"
        "<p style=\"margin:18px 0 0;font-size:13px;color:#64748b;\">This email was generated automatically by the IT Service Bot.</p>"
        "</div>"
        "</div>"
        "</body>"
        "</html>"
    )


def send_email_notification(
    recipient_email: str,
    subject: str,
    body: str,
    html_body: str | None = None,
) -> None:
    if not email_config_is_ready():
        missing = ", ".join(get_missing_smtp_settings())
        raise RuntimeError(
            f"SMTP settings are incomplete. Missing: {missing}. Set them in .env"
        )
    placeholders = get_placeholder_smtp_settings()
    if placeholders:
        raise RuntimeError(
            "SMTP settings still contain placeholder values: "
            + ", ".join(placeholders)
            + ". Replace them with your real Microsoft mail values in .env"
        )

    message = EmailMessage()
    message["Subject"] = subject
    message["From"] = SMTP_FROM_EMAIL
    message["To"] = recipient_email
    message.set_content(body)
    if html_body:
        message.add_alternative(html_body, subtype="html")

    with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=30) as server:
        server.ehlo()
        if SMTP_USE_TLS:
            server.starttls()
            server.ehlo()
        server.login(SMTP_USERNAME, SMTP_PASSWORD)
        server.send_message(message)


def send_email_with_attachment(recipient_email: str, subject: str, body: str, attachment_path: str) -> None:
    if not email_config_is_ready():
        missing = ", ".join(get_missing_smtp_settings())
        raise RuntimeError(
            f"SMTP settings are incomplete. Missing: {missing}. Set them in .env"
        )
    placeholders = get_placeholder_smtp_settings()
    if placeholders:
        raise RuntimeError(
            "SMTP settings still contain placeholder values: "
            + ", ".join(placeholders)
            + ". Replace them with your real Microsoft mail values in .env"
        )

    message = EmailMessage()
    message["Subject"] = subject
    message["From"] = SMTP_FROM_EMAIL
    message["To"] = recipient_email
    message.set_content(body)

    with open(attachment_path, "rb") as attachment_file:
        attachment_data = attachment_file.read()

    message.add_attachment(
        attachment_data,
        maintype="application",
        subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=os.path.basename(attachment_path),
    )

    with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=30) as server:
        server.ehlo()
        if SMTP_USE_TLS:
            server.starttls()
            server.ehlo()
        server.login(SMTP_USERNAME, SMTP_PASSWORD)
        server.send_message(message)


def now_local() -> datetime:
    return datetime.now(LOCAL_TZ)


def get_current_report_period(now: datetime | None = None) -> str:
    current = now or now_local()
    return current.strftime("%Y-%m")


def should_send_monthly_report(now: datetime | None = None) -> bool:
    current = now or now_local()
    if current.day != 1:
        return False
    if (current.hour, current.minute) < (MONTHLY_REPORT_SEND_HOUR, MONTHLY_REPORT_SEND_MINUTE):
        return False
    last_sent_period = get_setting("monthly_report_last_sent")
    return last_sent_period != get_current_report_period(current)


def seconds_until_next_monthly_run(now: datetime | None = None) -> float:
    current = now or now_local()
    target = current.replace(
        day=1,
        hour=MONTHLY_REPORT_SEND_HOUR,
        minute=MONTHLY_REPORT_SEND_MINUTE,
        second=0,
        microsecond=0,
    )
    if current >= target:
        if current.month == 12:
            target = target.replace(year=current.year + 1, month=1)
        else:
            target = target.replace(month=current.month + 1)
    return max((target - current).total_seconds(), 60.0)


async def run_monthly_report_job() -> None:
    if not MONTHLY_REPORT_ENABLED:
        logger.info("Monthly report skipped because MONTHLY_REPORT_ENABLED is false.")
        return
    if not MONTHLY_REPORT_RECIPIENT_EMAIL:
        logger.warning("Monthly report skipped because MONTHLY_REPORT_RECIPIENT_EMAIL is not configured.")
        return

    filename = export_to_excel()
    if filename is None:
        logger.error("Monthly report skipped because openpyxl is not installed.")
        return

    report_period = get_current_report_period()
    subject = f"Monthly IT Service Report - {report_period}"
    body = (
        f"Monthly IT service report for {report_period} is attached.\n\n"
        f"Generated at: {now_local().strftime('%Y-%m-%d %H:%M:%S %Z')}"
    )

    await asyncio.to_thread(
        send_email_with_attachment,
        MONTHLY_REPORT_RECIPIENT_EMAIL,
        subject,
        body,
        filename,
    )
    set_setting("monthly_report_last_sent", report_period)
    logger.info("Monthly report emailed to %s for period %s", MONTHLY_REPORT_RECIPIENT_EMAIL, report_period)


async def monthly_report_scheduler() -> None:
    if not MONTHLY_REPORT_ENABLED:
        logger.info("Monthly report scheduler is disabled.")
        return
    logger.info(
        "Monthly report scheduler started. Recipient=%s schedule=%02d:%02d %s",
        MONTHLY_REPORT_RECIPIENT_EMAIL,
        MONTHLY_REPORT_SEND_HOUR,
        MONTHLY_REPORT_SEND_MINUTE,
        APP_TIMEZONE,
    )
    while True:
        try:
            if should_send_monthly_report():
                await run_monthly_report_job()
                await asyncio.sleep(60)
                continue

            sleep_seconds = seconds_until_next_monthly_run()
            logger.info("Next monthly report check in %.0f seconds", sleep_seconds)
            await asyncio.sleep(sleep_seconds)
        except asyncio.CancelledError:
            logger.info("Monthly report scheduler stopped.")
            raise
        except Exception:
            logger.exception("Monthly report scheduler encountered an error.")
            await asyncio.sleep(300)


async def post_init(application: Application) -> None:
    if MONTHLY_REPORT_ENABLED:
        application.bot_data["monthly_report_task"] = asyncio.create_task(monthly_report_scheduler())


async def post_shutdown(application: Application) -> None:
    task = application.bot_data.get("monthly_report_task")
    if task:
        task.cancel()
        try:
            await task
        except asyncio.CancelledError:
            pass

# Commands (အမိန့်များ)
async def start(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Sends a technician-only welcome message."""
    email_line = (
        "Service email notifications are enabled."
        if EMAIL_DELIVERY_ENABLED
        else "Lite mode is enabled. Service emails and monthly auto emails are disabled."
    )
    report_line = (
        "Monthly Excel report ကို လစဉ် auto email ပို့မည်။"
        if MONTHLY_REPORT_ENABLED
        else "Monthly auto email report is disabled in this deployment."
    )
    await update.message.reply_text(
        "မင်္ဂလာပါ! ဒီ bot ကို Technician သီးသန့်အသုံးပြုပါသည်။\n"
        "Welcome! This bot is for technicians only.\n\n"
        "Commands:\n"
        "/newlog - ဝန်ဆောင်မှုမှတ်တမ်းအသစ် ဖန်တီးရန်\n"
        "/history - မှတ်တမ်းများ ကြည့်ရန်\n"
        "/export - CSV/Excel ဖိုင် export ထုတ်ရန်\n"
        f"{email_line}\n"
        f"{report_line}"
    )


async def prompt_for_service_type(chat_id: int, bot) -> None:
    keyboard = [
        [InlineKeyboardButton(option, callback_data=f"service_option_{index}")]
        for index, option in enumerate(SERVICE_OPTIONS)
    ]
    await bot.send_message(
        chat_id=chat_id,
        text="ကျေးဇူးပြု၍ Service Type ကို ရွေးချယ်ပါ။\nPlease choose the Service Type:",
        reply_markup=InlineKeyboardMarkup(keyboard),
    )

async def button(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Parses the CallbackQuery and updates the message text."""
    query = update.callback_query
    await query.answer()

    if query.data == "export_csv":
        await do_export_csv(update, context)
    elif query.data == "export_excel":
        await do_export_excel(update, context)
    elif query.data.startswith("service_option_"):
        option_index = int(query.data.rsplit("_", 1)[1])
        selected_service = SERVICE_OPTIONS[option_index]
        user_id = query.from_user.id

        if user_id not in user_data or user_data[user_id].get("state") != "waiting_service_done":
            await query.edit_message_text(
                "Service choice အခြေအနေ မတွေ့ပါ။ /newlog ဖြင့် ပြန်စတင်ပါ。\n"
                "Service choice state not found. Start again with /newlog."
            )
            return

        user_data[user_id]["service_done"] = selected_service
        user_data[user_id]["state"] = "waiting_parts_used"
        await query.edit_message_text(
            f"ရွေးချယ်ထားသော Service: {selected_service}\n"
            f"Selected service: {selected_service}"
        )
        await context.bot.send_message(
            chat_id=query.message.chat_id,
            text="ကျေးဇူးပြု၍ အသုံးပြုခဲ့သော အစိတ်အပိုင်းများ (Parts Used) ကို ထည့်သွင်းပါ။ (မရှိပါက 'None' ဟု ရိုက်ထည့်ပါ)\nPlease enter the Parts Used: (Type 'None' if none)"
        )

async def newlog(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Starts the process of creating a new service log."""
    user_data[update.effective_user.id] = {"state": "waiting_device_id"}
    await update.message.reply_text("ကျေးဇူးပြု၍ Device Name ကို ထည့်သွင်းပါ။\nPlease enter the Device Name:")

async def handle_message(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Handles messages based on the current state of the technician."""
    user_id = update.effective_user.id
    if user_id not in user_data:
        await update.message.reply_text("မသိသော command သို့မဟုတ် အခြေအနေ။ /start ဖြင့် စတင်ပါ။\nUnknown command or state. Use /start to begin.")
        return

    state = user_data[user_id]["state"]
    text = update.message.text

    if state == "waiting_device_id":
        user_data[user_id]["device_id"] = text
        user_data[user_id]["state"] = "waiting_user_name"
        await update.message.reply_text("ကျေးဇူးပြု၍ အသုံးပြုသူ၏အမည် (User Name) ကို ထည့်သွင်းပါ။\nPlease enter the User Name:")
    elif state == "waiting_user_name":
        user_data[user_id]["user_name"] = text
        if USER_EMAIL_REQUIRED:
            user_data[user_id]["state"] = "waiting_user_email"
            await update.message.reply_text("ကျေးဇူးပြု၍ အသုံးပြုသူ၏ Email Address ကို ထည့်သွင်းပါ။\nPlease enter the User Email Address:")
        else:
            user_data[user_id]["user_email"] = None
            user_data[user_id]["state"] = "waiting_service_done"
            await update.message.reply_text(
                "Lite mode ဖြစ်သောကြောင့် user email step ကို ကျော်သွားပါသည်。\n"
                "Lite mode is enabled, so the user email step is skipped."
            )
            await prompt_for_service_type(update.effective_chat.id, context.bot)
    elif state == "waiting_user_email":
        if not is_valid_email(text):
            await update.message.reply_text("Email format မမှန်ကန်ပါ။ ကျေးဇူးပြု၍ မှန်ကန်သော email address ကို ထည့်သွင်းပါ။\nInvalid email format. Please enter a valid email address.")
            return

        user_data[user_id]["user_email"] = text
        user_data[user_id]["state"] = "waiting_service_done"
        await prompt_for_service_type(update.effective_chat.id, context.bot)
    elif state == "waiting_service_done":
        await update.message.reply_text(
            "Service Done ကို button ဖြင့် ရွေးချယ်ပေးပါ။\n"
            "Please choose the service using the buttons above."
        )
    elif state == "waiting_parts_used":
        user_data[user_id]["parts_used"] = text if text.lower() != "none" else "N/A"
        await send_service_confirmation(update, context, user_id)
        del user_data[user_id]

async def send_service_confirmation(update: Update, context: ContextTypes.DEFAULT_TYPE, technician_user_id: int) -> None:
    """Sends the service notification email when enabled, otherwise stores the log only."""
    data = user_data[technician_user_id]
    technician_name = update.effective_user.full_name
    device_id = data["device_id"]
    user_name = data["user_name"]
    service_done = data["service_done"]
    parts_used = data["parts_used"]
    user_email = data["user_email"]
    service_date = now_local().strftime("%d-%m-%Y")

    log_id = add_service_log(
        technician_name=technician_name,
        device_id=device_id,
        user_name=user_name,
        user_email=user_email,
        service_done=service_done,
        parts_used=parts_used,
        technician_chat_id=update.effective_chat.id
    )

    if not EMAIL_DELIVERY_ENABLED:
        update_service_log_status(log_id, "Saved Locally (Lite Mode)")
        await update.message.reply_text(
            f"ဝန်ဆောင်မှုမှတ်တမ်း #{log_id} ကို local log အဖြစ် သိမ်းပြီးပါပြီ။\n"
            "Lite mode ဖြစ်သောကြောင့် email မပို့ပါ။\n"
            f"Service log #{log_id} saved locally. Email delivery is disabled in lite mode."
        )
        return

    subject = f"Service Slip - {device_id}"
    email_body = build_service_slip_text(
        user_name=user_name,
        device_id=device_id,
        technician_name=technician_name,
        service_done=service_done,
        parts_used=parts_used,
        service_date=service_date,
    )
    email_html = build_service_slip_html(
        user_name=user_name,
        device_id=device_id,
        technician_name=technician_name,
        service_done=service_done,
        parts_used=parts_used,
        service_date=service_date,
    )

    try:
        await asyncio.to_thread(send_email_notification, user_email, subject, email_body, email_html)
        await update.message.reply_text(
            f"ဝန်ဆောင်မှုမှတ်တမ်း #{log_id} ကို {user_name} ({user_email}) ထံ email ဖြင့် ပေးပို့ပြီးပါပြီ။\n"
            f"Service log #{log_id} sent to {user_name} ({user_email}) by email."
        )
    except Exception as exc:
        update_service_log_status(log_id, "Email Failed")
        logger.exception("Failed to send email for log %s", log_id)
        detail = describe_email_error(exc)
        await update.message.reply_text(
            "ဝန်ဆောင်မှု email ပို့ရာတွင် အမှားဖြစ်နေပါသည်။\n"
            f"Error: {exc}\n"
            f"Detail: {detail}\n"
            f"SMTP Host: {SMTP_HOST}\n"
            f"SMTP Port: {SMTP_PORT}\n"
            f"SMTP Username: {SMTP_USERNAME}\n"
            f"SMTP From: {SMTP_FROM_EMAIL}\n"
            f"TLS: {SMTP_USE_TLS}\n"
            "Please check .env values and restart the bot."
        )

async def history(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Displays past service logs for the technician."""
    chat_id = update.effective_chat.id

    logs = get_technician_service_logs(chat_id)
    title = "သင်၏ ဝန်ဆောင်မှုမှတ်တမ်းများ (Your Service Logs):\n"

    if not logs:
        await update.message.reply_text("မှတ်တမ်းများ မရှိသေးပါ။\nNo logs found.")
        return

    response_text = title
    for log in logs:
        response_text += (
            f"\nLog ID: {log['id']}\n"
            f"Device Name: {log['device_id']}\n"
            f"User Name: {log['user_name']}\n"
            f"User Email: {log['user_email'] or 'N/A'}\n"
            f"Technician: {log['technician_name']}\n"
            f"Service Done: {log['service_done']}\n"
            f"Parts Used: {log['parts_used']}\n"
            f"Status: {log['status']}\n"
            f"Date: {log['created_at']}\n"
            f"----------\n"
        )
    await update.message.reply_text(response_text)

# ============ Export Commands (Export အမိန့်များ) ============

async def export_command(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """/export command - CSV သို့မဟုတ် Excel ရွေးချယ်ရန်"""
    keyboard = [
        [InlineKeyboardButton("CSV ဖိုင် Export", callback_data="export_csv")],
        [InlineKeyboardButton("Excel ဖိုင် Export", callback_data="export_excel")],
    ]
    reply_markup = InlineKeyboardMarkup(keyboard)
    await update.message.reply_text(
        "Export ပုံစံကို ရွေးချယ်ပါ။\nChoose export format:",
        reply_markup=reply_markup,
    )

async def do_export_csv(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """CSV ဖိုင် export လုပ်ပြီး ပို့ပေးခြင်း"""
    query = update.callback_query
    await query.edit_message_text(text="CSV ဖိုင် ပြင်ဆင်နေပါသည်... ခဏစောင့်ပါ။\nPreparing CSV file... Please wait.")
    
    filename = export_to_csv()
    
    with open(filename, 'rb') as f:
        await context.bot.send_document(
            chat_id=query.message.chat_id,
            document=f,
            filename=os.path.basename(filename),
            caption="Service Log များ CSV ဖိုင်အဖြစ် export ပြီးပါပြီ။\nService logs exported as CSV file."
        )

async def do_export_excel(update: Update, context: ContextTypes.DEFAULT_TYPE) -> None:
    """Excel ဖိုင် export လုပ်ပြီး ပို့ပေးခြင်း"""
    query = update.callback_query
    await query.edit_message_text(text="Excel ဖိုင် ပြင်ဆင်နေပါသည်... ခဏစောင့်ပါ။\nPreparing Excel file... Please wait.")
    
    filename = export_to_excel()
    
    if filename is None:
        await context.bot.send_message(
            chat_id=query.message.chat_id,
            text="openpyxl library မရှိသေးပါ။ ကျေးဇူးပြု၍ install လုပ်ပါ:\npython -m pip install openpyxl\n\nCSV export ကို အစားထိုး သုံးနိုင်ပါသည်။"
        )
        return
    
    with open(filename, 'rb') as f:
        await context.bot.send_document(
            chat_id=query.message.chat_id,
            document=f,
            filename=os.path.basename(filename),
            caption="Service Log များ Excel ဖိုင်အဖြစ် export ပြီးပါပြီ။\nService logs exported as Excel file."
        )


def register_handlers(application: Application) -> None:
    # Register handlers (Handler များ မှတ်ပုံတင်ခြင်း)
    application.add_handler(CommandHandler("start", start))
    application.add_handler(CommandHandler("newlog", newlog))
    application.add_handler(CommandHandler("history", history))
    application.add_handler(CommandHandler("export", export_command))
    application.add_handler(CallbackQueryHandler(button))
    application.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_message))


def build_application() -> Application:
    application = (
        Application.builder()
        .token(TOKEN)
        .post_init(post_init)
        .post_shutdown(post_shutdown)
        .build()
    )
    register_handlers(application)
    return application


def validate_run_mode() -> str:
    if BOT_RUN_MODE not in {"polling", "webhook"}:
        raise RuntimeError("BOT_RUN_MODE must be either 'polling' or 'webhook'.")
    return BOT_RUN_MODE


def ensure_webhook_base_url() -> str:
    base_url = WEBHOOK_BASE_URL
    if not base_url:
        raise RuntimeError(
            "Webhook mode requires WEBHOOK_BASE_URL or Render's RENDER_EXTERNAL_URL to be set."
        )
    return base_url


def check_database_health() -> tuple[bool, str]:
    try:
        conn = get_db_connection()
        conn.execute("SELECT 1")
        conn.close()
    except Exception as exc:
        logger.exception("Health check database query failed.")
        return False, str(exc)
    return True, "ok"


async def healthcheck(request: Request) -> Response:
    application = request.app.state.telegram_application
    if not application.running:
        return PlainTextResponse("telegram application is starting", status_code=HTTPStatus.SERVICE_UNAVAILABLE)

    healthy, detail = check_database_health()
    if not healthy:
        return PlainTextResponse(
            f"database health check failed: {detail}",
            status_code=HTTPStatus.SERVICE_UNAVAILABLE,
        )

    return PlainTextResponse("ok", status_code=HTTPStatus.OK)


async def telegram_webhook(request: Request) -> Response:
    expected_secret = request.app.state.webhook_secret_token
    received_secret = request.headers.get("X-Telegram-Bot-Api-Secret-Token")
    if received_secret != expected_secret:
        return PlainTextResponse("invalid webhook secret", status_code=HTTPStatus.UNAUTHORIZED)

    application = request.app.state.telegram_application

    try:
        payload = await request.json()
        update = Update.de_json(payload, application.bot)
    except Exception:
        logger.exception("Failed to parse incoming Telegram webhook payload.")
        return PlainTextResponse("invalid payload", status_code=HTTPStatus.BAD_REQUEST)

    await application.update_queue.put(update)
    return Response(status_code=HTTPStatus.OK)


def build_web_app(application: Application) -> Starlette:
    webhook_url = build_webhook_url(ensure_webhook_base_url(), WEBHOOK_PATH)

    @asynccontextmanager
    async def lifespan(starlette_app: Starlette):
        await application.initialize()
        await application.start()
        await post_init(application)
        await application.bot.set_webhook(
            url=webhook_url,
            allowed_updates=Update.ALL_TYPES,
            secret_token=WEBHOOK_SECRET_TOKEN,
            drop_pending_updates=DROP_PENDING_UPDATES,
        )
        logger.info("Webhook configured at %s", webhook_url)

        try:
            yield
        finally:
            logger.info("Stopping webhook application.")
            await post_shutdown(application)
            await application.stop()
            await application.shutdown()

    app = Starlette(
        routes=[
            Route(WEBHOOK_PATH, telegram_webhook, methods=["POST"]),
            Route(HEALTHCHECK_PATH, healthcheck, methods=["GET"]),
        ],
        lifespan=lifespan,
    )
    app.state.telegram_application = application
    app.state.webhook_secret_token = WEBHOOK_SECRET_TOKEN
    return app


def run_polling_mode(application: Application) -> None:
    logger.info("Starting bot in polling mode.")
    application.run_polling(allowed_updates=Update.ALL_TYPES)


def run_webhook_mode(application: Application) -> None:
    web_app = build_web_app(application)
    logger.info("Starting bot in webhook mode on %s:%s", WEBHOOK_LISTEN_HOST, PORT)
    uvicorn.run(web_app, host=WEBHOOK_LISTEN_HOST, port=PORT, use_colors=False)


def main() -> None:
    """Start the bot."""
    run_mode = validate_run_mode()
    init_db()
    logger.info("Using database at %s", DB_PATH)
    logger.info("Using export directory at %s", EXPORT_DIR)
    logger.info("Bot run mode: %s", run_mode)
    logger.info("Free plan mode: %s", FREE_PLAN_MODE)
    logger.info("Email delivery enabled: %s", EMAIL_DELIVERY_ENABLED)
    logger.info("Monthly report enabled: %s", MONTHLY_REPORT_ENABLED)
    if WEBHOOK_BASE_URL:
        logger.info("Webhook base URL: %s", WEBHOOK_BASE_URL)

    application = build_application()

    if run_mode == "webhook":
        run_webhook_mode(application)
        return

    run_polling_mode(application)


if __name__ == "__main__":
    main()
